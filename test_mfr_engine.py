from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from benchmark_engine import read_order_file
from mfr_engine import (
    MFR_METRIC_COLUMNS,
    process_shopee_mfr,
    process_tik_family_mfr,
    process_lazada_mfr,
    process_dataframe_mfr,
    calculate_mfr,
    validate_mfr_file,
    create_mfr_excel_report,
)

print("=== STARTING MFR ENGINE TESTS ===")

# -------------------------------------------------------------
# TEST 1: SHOPEE MFR
# -------------------------------------------------------------
sho_df = pd.DataFrame({
    "No. Pesanan": ["SHO-001", "SHO-002", "SHO-003", "SHO-003", "SHO-004"],
    "Status Pesanan": ["Selesai", "Sedang Dikemas", "Dikirim", "Dikirim", "Batal"],
    "Waktu Pesanan Dibuat": [
        "2026-09-01 10:00",
        "2026-09-05 11:00",
        "2026-09-10 12:00",
        "2026-09-10 12:00",
        "2026-09-15 13:00",
    ],
    "Subtotal Pesanan": ["100.000", "200.000", "150.000", "50.000", "300.000"],
    "Voucher Ditanggung Penjual": ["10.000", "20.000", "15.000", "15.000", "50.000"],
    "Diskon Dari Shopee": ["5.000", "10.000", "5.000", "5.000", "20.000"],
})

res_sho = process_shopee_mfr(sho_df)
m_sho = res_sho["monthly"]
assert len(m_sho) == 1
r_sho = m_sho.iloc[0]

# Non-batal rows: SHO-001 (100k), SHO-002 (200k), SHO-003 (150k + 50k = 200k).
# Total GMV non-batal = 500k.
# Deduplicated Seller Rebate: SHO-001 (10k) + SHO-002 (20k) + SHO-003 (max 15k) = 45k.
# Diskon Shopee non-batal: 5k + 10k + 5k + 5k = 25k.
# Net GMV = 500k - 45k + 25k = 480k.
assert r_sho["gmv_all"] == 800000.0
assert r_sho["gmv_ex_cancel"] == 500000.0
assert r_sho["seller_rebate"] == 45000.0
assert r_sho["shopee_discount"] == 25000.0
assert r_sho["platform_discount"] == 0.0
assert r_sho["net_gmv"] == 480000.0
assert r_sho["source_rows"] == 5
assert r_sho["included_rows"] == 4
assert r_sho["excluded_rows"] == 1
sho_detail = res_sho["order_detail"]
sho_003_detail = sho_detail[sho_detail["order_id"] == "SHO-003"]
assert sho_003_detail["sequence"].tolist() == [1, 2]
assert sho_003_detail["seller_rebate"].tolist() == [15000.0, 0.0]
assert sho_003_detail["shopee_discount"].tolist() == [10000.0, 0.0]
assert sho_detail["seller_rebate"].sum() == r_sho["seller_rebate"]
assert sho_detail["shopee_discount"].sum() == r_sho["shopee_discount"]
assert sho_detail["platform_discount"].sum() == 0
print("✅ Test 1 (Shopee MFR) PASSED")

# -------------------------------------------------------------
# TEST 2: TIKTOK MFR
# -------------------------------------------------------------
tik_df = pd.DataFrame({
    "Order ID": ["T-01", "T-02", "T-03", "T-04", "T-05"],
    "Order Sub Status": ["Selesai", "Dikirim", "Dalam Transit", "Dibatalkan", ""],
    "Order Status": ["Selesai", "Dikirim", "Dalam Transit", "Dibatalkan", "Dikirim"],
    "Cancelation/Return Type": ["", "", "", "Cancel", ""],
    "Created Time": [
        "01/09/2026 10:00:00",
        "05/09/2026 11:00:00",
        "10/09/2026 12:00:00",
        "15/09/2026 13:00:00",
        "20/09/2026 14:00:00",
    ],
    "SKU Subtotal After Discount": [100000, 150000, 80000, 200000, 50000],
    "SKU Platform Discount": [10000, 15000, 8000, 20000, 5000],
    "Purchase Channel": ["TikTok"] * 5,
})

res_tik = process_tik_family_mfr(tik_df, "TIK")
m_tik = res_tik["monthly"]
assert len(m_tik) == 1
r_tik = m_tik.iloc[0]

# Blank substatus T-05 stays included because it does not match a cancel rule.
# T-04 is excluded by the full cancel rule.
assert r_tik["gmv_ex_cancel"] == 380000.0
assert r_tik["platform_discount"] == 38000.0
assert r_tik["net_gmv"] == 418000.0
assert r_tik["included_rows"] == 4
assert r_tik["excluded_rows"] == 1
assert len(res_tik["warnings"]) >= 1  # warning for blank status
assert res_tik["order_detail"]["platform_discount"].sum() == r_tik[
    "platform_discount"
]
assert res_tik["order_detail"]["seller_rebate"].sum() == 0
assert res_tik["order_detail"]["shopee_discount"].sum() == 0
print("✅ Test 2 (TikTok MFR) PASSED")

# -------------------------------------------------------------
# TEST 3: TOKOPEDIA MFR
# -------------------------------------------------------------
tok_df = pd.DataFrame({
    "Order ID": ["TOK-01", "TOK-02"],
    "Order Sub Status": ["Selesai", "Dalam Pengiriman"],
    "Order Status": ["Selesai", "Dalam Pengiriman"],
    "Cancelation/Return Type": ["", ""],
    "Created Time": ["02/09/2026 09:00:00", "03/09/2026 10:00:00"],
    "SKU Subtotal After Discount": [120000, 80000],
    "SKU Platform Discount": [12000, 8000],
    "Purchase Channel": ["Tokopedia", "Tokopedia"],
})
res_tok = process_tik_family_mfr(tok_df, "TOK")
r_tok = res_tok["monthly"].iloc[0]
assert r_tok["net_gmv"] == 220000.0
assert r_tok["included_rows"] == 2
assert r_tok["excluded_rows"] == 0
assert res_tok["order_detail"]["platform_discount"].sum() == r_tok[
    "platform_discount"
]
print("✅ Test 3 (Tokopedia MFR) PASSED")

# -------------------------------------------------------------
# TEST 4: LAZADA MFR
# -------------------------------------------------------------
laz_df = pd.DataFrame({
    "orderItemId": ["LAZ-01", "LAZ-02", "LAZ-03"],
    "createTime": ["01/09/2026 08:00:00", "05/09/2026 09:00:00", "10/09/2026 10:00:00"],
    "unitPrice": [90000, 110000, 300000],
    "platformDiscountTotal": [10000, 15000, 50000],
    "status": ["delivered", "shipped", "canceled"],
})
res_laz = process_lazada_mfr(laz_df)
r_laz = res_laz["monthly"].iloc[0]
# Included: LAZ-01 (90k+10k=100k), LAZ-02 (110k+15k=125k) = 225k.
# Excluded: LAZ-03 (canceled)
assert r_laz["net_gmv"] == 225000.0
assert r_laz["included_rows"] == 2
assert r_laz["excluded_rows"] == 1
assert res_laz["order_detail"]["platform_discount"].sum() == r_laz[
    "platform_discount"
]
print("✅ Test 4 (Lazada MFR) PASSED")

# -------------------------------------------------------------
# TEST 5: MULTI-MONTH DATA & 1 MONTH SELECTION
# -------------------------------------------------------------
# Add August data to Shopee to test multi-month isolation
sho_multi_df = pd.concat([
    sho_df,
    pd.DataFrame({
        "No. Pesanan": ["SHO-AUG-1"],
        "Status Pesanan": ["Selesai"],
        "Waktu Pesanan Dibuat": ["2026-08-15 10:00"],
        "Subtotal Pesanan": ["500.000"],
        "Voucher Ditanggung Penjual": ["0"],
        "Diskon Dari Shopee": ["0"],
    })
], ignore_index=True)

res_sho_multi = process_shopee_mfr(sho_multi_df)
all_results = [res_sho_multi, res_tik, res_tok, res_laz]

# Calculate MFR for September 2026
mfr_final = calculate_mfr(all_results, selected_month="2026-09-01")
assert mfr_final["month_label"] == "September 2026"
# SHO (480k) + TIK (418k) + TOK (220k) + LAZ (225k) = 1,343,000
expected_total = 480000 + 418000 + 220000 + 225000
assert mfr_final["mfr_total"] == expected_total, f"Expected {expected_total}, got {mfr_final['mfr_total']}"
assert mfr_final["marketplace_count"] == 4
# Check that August is completely excluded from September MFR
summary = mfr_final["marketplace_summary"]
sho_sum_omzet = summary[summary["marketplace"] == "SHO"]["net_gmv"].iloc[0]
assert sho_sum_omzet == 480000.0, f"Shopee MFR in Sept should be 480k, got {sho_sum_omzet}"
print(f"✅ Test 5 (Multi-month isolation & Total MFR = Rp {mfr_final['mfr_total']:,.0f}) PASSED")

# -------------------------------------------------------------
# TEST 6: STATUS BREAKDOWN
# -------------------------------------------------------------
sb = mfr_final["status_breakdown"]
assert not sb.empty
assert set(sb["marketplace"]) == {"SHO", "TIK", "TOK", "LAZ"}
assert set(sb["treatment"]) == {"Included", "Excluded"}
# Check that all canceled/dibatalkan are Excluded
cancels = sb[sb["status"].str.upper().str.contains("BATAL|CANCEL")]
assert (cancels["treatment"] == "Excluded").all()
print("✅ Test 6 (Status Breakdown & Treatment) PASSED")

# -------------------------------------------------------------
# TEST 7: ZERO MISTAKE AUDIT
# -------------------------------------------------------------
audit = mfr_final["audit"]
fail_audits = [a for a in audit if a["Status"] == "FAIL"]
assert len(fail_audits) == 0, f"Audit failed: {fail_audits}"
print("✅ Test 7 (Zero Mistake Audit) PASSED")

# -------------------------------------------------------------
# TEST 8: EXCEL REPORT GENERATION
# -------------------------------------------------------------
excel_io = create_mfr_excel_report(mfr_final)
excel_bytes = excel_io.getvalue()
assert len(excel_bytes) > 1000
print(f"✅ Test 8 (Excel Report: {len(excel_bytes)} bytes) PASSED")

# -------------------------------------------------------------
# TEST 9: SHOPEE NUMERIC SOURCE UNIT NORMALIZATION
# -------------------------------------------------------------
sho_numeric_df = pd.DataFrame({
    "No. Pesanan": ["SHO-NUM-1"],
    "Status Pesanan": ["Selesai"],
    "Waktu Pesanan Dibuat": ["2026-09-20 10:00"],
    "Subtotal Pesanan": [190.0],
    "Voucher Ditanggung Penjual": [8700],
    "Diskon Dari Shopee": [9.5],
})
numeric_row = process_shopee_mfr(sho_numeric_df)["monthly"].iloc[0]
assert numeric_row["gmv_all"] == 190000.0
assert numeric_row["gmv_ex_cancel"] == 190000.0
assert numeric_row["seller_rebate"] == 8700.0
assert numeric_row["shopee_discount"] == 9500.0
assert numeric_row["net_gmv"] == 190800.0
print("✅ Test 9 (Shopee numeric source normalization) PASSED")

# -------------------------------------------------------------
# TEST 10: GOJI-M AUGUST 2026 EXACT REGRESSION
# -------------------------------------------------------------
data_dir = Path(__file__).resolve().parent / "contoh data pesanan GOJI MFR"
real_results = []
for path in sorted(data_dir.glob("Daftar Pesanan*.xlsx")):
    if "SHO " in path.name:
        marketplace = "SHO"
    elif path.stem.endswith("_TOK"):
        marketplace = "TOK"
    else:
        marketplace = "TIK"

    with path.open("rb") as source:
        real_df, detected, _ = read_order_file(
            source,
            filename=path.name,
            marketplace_hint=(
                marketplace if marketplace in {"TIK", "TOK"} else None
            ),
        )
    assert detected == marketplace
    real_results.append(
        process_dataframe_mfr(real_df, marketplace, filename=path.name)
    )

goji = calculate_mfr(real_results, selected_month="2026-08-01")
expected = {
    "SHO": {
        "gmv_all": 212008536,
        "gmv_ex_cancel": 172509867,
        "seller_rebate": 6433295,
        "shopee_discount": 9500,
        "platform_discount": 0,
        "net_gmv": 166086072,
    },
    "TIK": {
        "gmv_all": 398085385,
        "gmv_ex_cancel": 288373190,
        "seller_rebate": 0,
        "shopee_discount": 0,
        "platform_discount": 13697681,
        "net_gmv": 302070871,
    },
    "TOK": {
        "gmv_all": 4615435,
        "gmv_ex_cancel": 3752228,
        "seller_rebate": 0,
        "shopee_discount": 0,
        "platform_discount": 287907,
        "net_gmv": 4040135,
    },
}
for _, actual in goji["marketplace_detail"].iterrows():
    marketplace = actual["marketplace"]
    for metric in MFR_METRIC_COLUMNS:
        assert actual[metric] == expected[marketplace][metric], (
            marketplace,
            metric,
            actual[metric],
            expected[marketplace][metric],
        )
assert not [row for row in goji["audit"] if row["Status"] == "FAIL"]
assert goji["mfr_total"] == 472197078
print("✅ Test 10 (GOJI-M August 2026 exact regression) PASSED")

# -------------------------------------------------------------
# TEST 11: DAFTAR PESANAN REFERENCE + RECONCILIATION + EXPORT
# -------------------------------------------------------------
order_detail = goji["order_detail"]
assert len(order_detail) == 2815
assert order_detail["paid_amount"].dtype.kind in {"f", "i", "u"}
for component in ["seller_rebate", "shopee_discount", "platform_discount"]:
    assert order_detail[component].dtype.kind in {"f", "i", "u"}
assert set(order_detail["marketplace"]) == {"SHO", "TIK", "TOK"}
assert order_detail["order_date"].dt.to_period("M").astype(str).eq("2026-08").all()

expected_component_totals = {
    "SHO": (6433295, 9500, 0),
    "TIK": (0, 0, 13697681),
    "TOK": (0, 0, 287907),
}
for marketplace, component_totals in expected_component_totals.items():
    rows = order_detail[order_detail["marketplace"] == marketplace]
    assert tuple(
        rows[column].sum()
        for column in ["seller_rebate", "shopee_discount", "platform_discount"]
    ) == component_totals

shopee_nonfirst = order_detail[
    order_detail["marketplace"].eq("SHO") & order_detail["sequence"].ne(1)
]
assert shopee_nonfirst[["seller_rebate", "shopee_discount"]].eq(0).all().all()

expected_detail_totals = {
    "SHO": (973, 166086072),
    "TIK": (1819, 302070871),
    "TOK": (23, 4040135),
}
for marketplace, (expected_rows, expected_paid) in expected_detail_totals.items():
    marketplace_rows = order_detail[order_detail["marketplace"] == marketplace]
    assert len(marketplace_rows) == expected_rows
    assert marketplace_rows["paid_amount"].sum() == expected_paid

sequence_examples = {
    "26080169SDJT1B": (1, 1),
    "260805HFTWQKDB": (3, 3),
    # Five distinct SKUs, each Qty=2 in the reference, produce ten unit rows.
    "260808QT0K2Q1S": (5, 10),
}
for order_id, (expected_skus, expected_rows) in sequence_examples.items():
    rows = order_detail[order_detail["order_id"] == order_id]
    assert rows["sku"].nunique() == expected_skus
    assert len(rows) == expected_rows
    assert rows["sequence"].tolist() == list(range(1, expected_rows + 1))

reference_path = data_dir / (
    "ID Monthly Financial Report (MFR) MOSO - SHO TIK - GOJI-M "
    "(10PFR) - AUG 2026 - Update.xlsx"
)
reference_workbook = load_workbook(reference_path, read_only=True, data_only=True)
reference_sheet = reference_workbook["Daftar Pesanan"]
reference_rows = []
for row in reference_sheet.iter_rows(min_row=1, max_col=10, values_only=True):
    if all(value is None or value == "" for value in row):
        if reference_rows:
            break
        continue
    reference_rows.append(row)
reference_workbook.close()
reference = pd.DataFrame(reference_rows[1:], columns=reference_rows[0])
reference["Tanggal"] = pd.to_datetime(
    reference["Tanggal"], format="%d %b %Y"
)

program_compare = order_detail.copy()
program_compare["reference_marketplace"] = program_compare["marketplace"].map(
    {"SHO": "SHO/GOJI-M", "TIK": "TIK/GOJI-M", "TOK": "TIK/GOJI-M"}
)
comparison = reference.merge(
    program_compare,
    left_on=["Marketplace-Toko", "Nomor Pesanan", "SKU", "Urutan"],
    right_on=["reference_marketplace", "order_id", "sku", "sequence"],
    how="outer",
    indicator=True,
)
assert comparison["_merge"].eq("both").all()
assert comparison["Tanggal"].eq(comparison["order_date"]).all()
assert comparison["Status Pesanan"].fillna("").eq(
    comparison["order_status"].fillna("")
).all()
assert comparison["Alasan Pembatalan"].fillna("").eq(
    comparison["cancellation_reason"].fillna("")
).all()
# Eight unit rows intentionally differ by Rp1 to absorb reference rounding
# residuals and make every marketplace reconcile to Net GMV exactly.
paid_difference = comparison["paid_amount"] - comparison["Harga Terbayarkan"]
assert paid_difference.abs().le(1).all()
assert paid_difference.ne(0).sum() == 8

goji_excel = create_mfr_excel_report(goji)
export_workbook = load_workbook(goji_excel, read_only=True, data_only=True)
assert "Daftar Pesanan" in export_workbook.sheetnames
export_sheet = export_workbook["Daftar Pesanan"]
expected_headers = [
    "Tanggal",
    "Marketplace",
    "Nomor Pesanan",
    "SKU",
    "Urutan",
    "Harga Terbayarkan",
    "Seller Rebate",
    "Diskon Shopee",
    "Platform Discount",
    "Status Pesanan",
    "Alasan Pembatalan",
]
assert [cell.value for cell in next(export_sheet.iter_rows(max_row=1))] == expected_headers
assert export_sheet.max_row == 2816
assert isinstance(export_sheet["A2"].value, pd.Timestamp) or hasattr(
    export_sheet["A2"].value, "year"
)
assert export_sheet["A2"].number_format == "d mmm yyyy"
assert export_sheet["C2"].number_format == "@"
assert export_sheet["D2"].number_format == "@"
assert isinstance(export_sheet["E2"].value, int)
assert isinstance(export_sheet["F2"].value, (int, float))
for coordinate in ["F2", "G2", "H2", "I2"]:
    assert export_sheet[coordinate].number_format == '"Rp "#,##0'
export_data = list(export_sheet.iter_rows(min_row=2, max_col=11, values_only=True))
assert sum(row[5] for row in export_data) == 472197078
assert sum(row[6] for row in export_data) == 6433295
assert sum(row[7] for row in export_data) == 9500
assert sum(row[8] for row in export_data) == 13985588
assert all(isinstance(row[index], (int, float)) for row in export_data for index in range(5, 9))
assert export_data[0][4] == 1
export_workbook.close()
print("✅ Test 11 (Daftar Pesanan reference, reconciliation, export) PASSED")

print("=== ALL MFR TESTS PASSED SUCCESSFULLY! ===")
