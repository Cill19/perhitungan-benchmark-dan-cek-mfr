import pandas as pd
import openpyxl
from pathlib import Path


# ============================================================
# CONFIG
# ============================================================

FILE_PATH = "data pesanan sho th.xlsx"


# ============================================================
# SHOW EXCEL STRUCTURE
# ============================================================

def diagnose_excel(file_path):

    path = Path(file_path)

    if not path.exists():
        print(
            f"File tidak ditemukan: {file_path}"
        )
        return


    print("=" * 100)
    print(
        f"FILE: {path.name}"
    )
    print("=" * 100)


    # --------------------------------------------------------
    # OPENPYXL CHECK
    # --------------------------------------------------------

    print("\n[OPENPYXL]")

    try:

        wb = openpyxl.load_workbook(
            file_path,
            read_only=True,
            data_only=True
        )

        print(
            "Sheets:",
            wb.sheetnames
        )


        for ws in wb.worksheets:

            print(
                f"\nSheet: {ws.title}"
            )

            print(
                "max_row:",
                ws.max_row,
                "| max_column:",
                ws.max_column
            )


            for row in ws.iter_rows(
                min_row=1,
                max_row=min(5, ws.max_row),
                values_only=True
            ):

                print(
                    list(row)[:20]
                )


    except Exception as e:

        print(
            "OPENPYXL ERROR:",
            e
        )



    # --------------------------------------------------------
    # PANDAS CHECK
    # --------------------------------------------------------

    print("\n[PANDAS]")

    try:

        excel = pd.ExcelFile(
            file_path
        )


        for sheet in excel.sheet_names:

            print(
                f"\nSheet: {sheet}"
            )


            df = pd.read_excel(
                file_path,
                sheet_name=sheet,
                header=None,
                nrows=5
            )


            print(
                "Shape preview:",
                df.shape
            )


            for index, row in df.iterrows():

                print(
                    f"row {index+1}:",
                    row.dropna()
                    .tolist()[:20]
                )


    except Exception as e:

        print(
            "PANDAS ERROR:",
            e
        )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    diagnose_excel(
        FILE_PATH
    )