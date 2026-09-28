"""
MFR Engine - Perhitungan MFR (Monthly Final Revenue / Net GMV 1 Bulan).

Prinsip Bisnis MFR:
- Menghitung Net GMV untuk SATU calendar month tertentu.
- SEMUA pesanan pada bulan tersebut dianggap masuk omzet SELAMA PESANAN BELUM / TIDAK DIBATALKAN.
- Status aktif (Diproses, Sedang Dikemas, Dikirim, Dalam Transit, Selesai, Terkirim, dll.) tetap INCLUDED.
- Hanya status final BATAL / CANCELED yang EXCLUDED.
- Filter pembatalan TikTok/Tokopedia menggunakan 3 layer:
  1. Order Sub Status == DIBATALKAN
  2. Order Status mengandung BATAL|CANCEL|CANCELED|CANCELLED
  3. Cancelation/Return Type mengandung CANCEL
- Pemisahan marketplace TikTok vs Tokopedia didasarkan pada kolom Purchase Channel.
- TIDAK ADA AVERAGE MULTI-BULAN.
"""

from __future__ import annotations

import io
from numbers import Number
from typing import Any, Callable, Iterable

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from benchmark_engine import (
    BenchmarkInputError,
    _month_start,
    _nonempty_invalid_numeric_count,
    _normalized_id_series,
    _parse_datetime,
    _require_columns,
    _resolve_column,
    _to_idr_number,
)


MFR_METRIC_COLUMNS = [
    "gmv_all",
    "gmv_ex_cancel",
    "seller_rebate",
    "shopee_discount",
    "platform_discount",
    "net_gmv",
]

MFR_MONTHLY_COLUMNS = [
    "marketplace",
    "month",
    *MFR_METRIC_COLUMNS,
    "source_rows",
    "included_rows",
    "excluded_rows",
]

MFR_STATUS_COLUMNS = [
    "marketplace",
    "month",
    "status",
    "rows",
    "net_gmv",
    "treatment",
]

MFR_ORDER_DETAIL_COLUMNS = [
    "order_date",
    "marketplace",
    "order_id",
    "sku",
    "sequence",
    "paid_amount",
    "seller_rebate",
    "shopee_discount",
    "platform_discount",
    "order_status",
    "cancellation_reason",
    "is_mfr_included",
]

MFR_ORDER_DETAIL_COMPONENT_COLUMNS = [
    "seller_rebate",
    "shopee_discount",
    "platform_discount",
]

MFR_DISPLAY_LABELS = {
    "marketplace": "Marketplace",
    "gmv_all": "GMV ALL",
    "gmv_ex_cancel": "GMV Exc Batal",
    "seller_rebate": "Seller Rebate",
    "shopee_discount": "Diskon Shopee",
    "platform_discount": "Platform Discount",
    "net_gmv": "Net GMV",
    "source_rows": "Rows Total",
    "included_rows": "Rows Included",
    "excluded_rows": "Rows Excluded",
    "status": "Status Pesanan",
    "rows": "Rows",
    "treatment": "Treatment",
}

MFR_ORDER_DETAIL_DISPLAY_LABELS = {
    "order_date": "Tanggal",
    "marketplace": "Marketplace",
    "order_id": "Nomor Pesanan",
    "sku": "SKU",
    "sequence": "Urutan",
    "paid_amount": "Harga Terbayarkan",
    "seller_rebate": "Seller Rebate",
    "shopee_discount": "Diskon Shopee",
    "platform_discount": "Platform Discount",
    "order_status": "Status Pesanan",
    "cancellation_reason": "Alasan Pembatalan",
}


def _require_internal_schema(
    df: pd.DataFrame,
    required: Iterable[str],
    context: str,
) -> None:
    """Fail with a useful contract error instead of an opaque pandas KeyError."""
    required_columns = list(required)
    missing = [column for column in required_columns if column not in df.columns]
    if missing:
        raise BenchmarkInputError(
            f"Schema MFR tidak valid pada {context}. "
            f"Kolom hilang: {missing}. Kolom tersedia: {df.columns.tolist()}"
        )


def _normalize_shopee_amount(
    series: pd.Series,
    *,
    multiply_numeric: Callable[[pd.Series], pd.Series],
) -> pd.Series:
    """
    Normalize Shopee IDR without applying the numeric-unit rule to text currency.

    A numeric cell such as ``190`` can represent Rp190.000 and is eligible for
    the supplied multiplier rule. A text value such as ``"190.000"`` is
    already a formatted currency amount and must only have punctuation removed.
    """
    cleaned = _to_idr_number(series).astype(float)
    numeric_mask = series.map(
        lambda value: isinstance(value, Number)
        and not isinstance(value, (bool, np.bool_))
        and not pd.isna(value)
    )
    numeric_values = pd.to_numeric(series.where(numeric_mask), errors="coerce")
    multiplier_mask = numeric_mask & multiply_numeric(numeric_values).fillna(False)
    cleaned.loc[numeric_mask] = numeric_values.loc[numeric_mask].astype(float)
    cleaned.loc[multiplier_mask] = (
        numeric_values.loc[multiplier_mask].astype(float) * 1000.0
    )
    return cleaned


def _empty_order_detail() -> pd.DataFrame:
    return pd.DataFrame(columns=MFR_ORDER_DETAIL_COLUMNS)


def _standardize_order_detail_schema(
    df: pd.DataFrame,
    context: str,
) -> pd.DataFrame:
    """Return one canonical, calculation-safe Daftar Pesanan schema."""
    standardized = df.copy()

    # Component columns are shared by every marketplace. Missing/non-applicable
    # components must be numeric zero so concat cannot create sparse/object
    # columns and downstream calculations never depend on display labels.
    for column in MFR_ORDER_DETAIL_COMPONENT_COLUMNS:
        if column not in standardized.columns:
            standardized[column] = 0.0

    _require_internal_schema(
        standardized,
        MFR_ORDER_DETAIL_COLUMNS,
        context,
    )

    standardized["paid_amount"] = pd.to_numeric(
        standardized["paid_amount"], errors="raise"
    )
    for column in MFR_ORDER_DETAIL_COMPONENT_COLUMNS:
        standardized[column] = (
            pd.to_numeric(standardized[column], errors="raise")
            .fillna(0.0)
            .astype(float)
        )

    # Enforce non-applicable marketplace components as numeric zero.
    marketplace = standardized["marketplace"].astype("string").str.upper()
    standardized.loc[marketplace.eq("SHO"), "platform_discount"] = 0.0
    standardized.loc[
        marketplace.isin(["TIK", "TOK", "LAZ"]),
        ["seller_rebate", "shopee_discount"],
    ] = 0.0

    return standardized[MFR_ORDER_DETAIL_COLUMNS].copy()


def _quantity_or_one(series: pd.Series) -> pd.Series:
    """Return positive integer quantities; invalid/zero values safely mean one."""
    quantity = pd.to_numeric(series, errors="coerce").fillna(1)
    quantity = quantity.where(quantity.gt(0), 1).round().astype(int)
    return quantity


def _round_half_away_from_zero(values: pd.Series) -> pd.Series:
    """Excel-compatible whole-IDR rounding for positive and negative values."""
    numeric = pd.to_numeric(values, errors="coerce").fillna(0.0).astype(float)
    rounded = np.where(
        numeric.ge(0),
        np.floor(numeric + 0.5),
        np.ceil(numeric - 0.5),
    )
    return pd.Series(rounded, index=values.index, dtype=float)


def _expand_and_reconcile_paid_amount(
    lines: pd.DataFrame,
    *,
    group_columns: list[str],
    exact_unit_amount: pd.Series,
    target_amount: pd.Series,
) -> pd.DataFrame:
    """
    Expand source lines by quantity and reconcile rounding at order level.

    The reference workbook repeats quantity units and rounds each unit to whole
    IDR. The final unit receives any rounding residual so detail always sums
    exactly to the validated MFR Net GMV.
    """
    if lines.empty:
        return lines.copy()

    source = lines.copy()
    source["_exact_unit_amount"] = exact_unit_amount.astype(float)
    source["_target_amount"] = target_amount.astype(float)
    expanded = source.loc[source.index.repeat(source["quantity"])].copy()
    expanded["_unit_position"] = expanded.groupby(
        "_source_row", sort=False
    ).cumcount()
    expanded["paid_amount"] = _round_half_away_from_zero(
        expanded["_exact_unit_amount"]
    )

    group_paid = expanded.groupby(group_columns, sort=False)["paid_amount"].transform(
        "sum"
    )
    residual = expanded["_target_amount"] - group_paid
    is_last = expanded.groupby(group_columns, sort=False).cumcount(ascending=False).eq(0)
    expanded.loc[is_last, "paid_amount"] += residual.loc[is_last]
    return expanded


def _audit_row(
    status: str,
    check: str,
    detail: str,
    *,
    actual: Any = "",
    expected: Any = "",
) -> dict[str, Any]:
    return {
        "Status": status,
        "Check": check,
        "Detail": detail,
        "Actual": actual,
        "Expected": expected,
    }


# ============================================================
# 1. SHOPEE MFR
# ============================================================

def build_shopee_order_detail(df: pd.DataFrame) -> pd.DataFrame:
    """Build eligible Shopee SKU/unit detail using the reference allocation."""
    required = {
        "No. Pesanan": (),
        "Status Pesanan": ("Order Status", "Status"),
        "Waktu Pesanan Dibuat": (),
        "Subtotal Pesanan": (),
        "Voucher Ditanggung Penjual": (),
        "Diskon Dari Shopee": (),
    }
    cols = _require_columns(df, "SHO", required)
    sku_col = _resolve_column(
        df,
        "Nomor Referensi SKU",
        ["Seller SKU", "SKU Penjual", "SKU"],
    )
    quantity_col = _resolve_column(df, "Jumlah", ["Quantity", "Qty"])
    unit_price_col = _resolve_column(
        df,
        "Harga Setelah Diskon",
        ["Discounted Price", "Harga Produk Setelah Diskon"],
    )
    reason_col = _resolve_column(
        df,
        "Alasan Pembatalan",
        ["Cancellation Reason", "Cancel Reason"],
    )

    order_id = _normalized_id_series(df[cols["No. Pesanan"]])
    status = df[cols["Status Pesanan"]].astype("string").fillna("").str.strip()
    created_at = _parse_datetime(df[cols["Waktu Pesanan Dibuat"]], dayfirst=False)
    subtotal = _normalize_shopee_amount(
        df[cols["Subtotal Pesanan"]],
        multiply_numeric=lambda values: values.abs().lt(10000) & values.ne(0),
    )
    seller_rebate = _to_idr_number(df[cols["Voucher Ditanggung Penjual"]])
    shopee_discount = _normalize_shopee_amount(
        df[cols["Diskon Dari Shopee"]],
        multiply_numeric=lambda values: values.gt(0) & values.lt(100),
    )
    if unit_price_col is not None:
        unit_base = _normalize_shopee_amount(
            df[unit_price_col],
            multiply_numeric=lambda values: values.abs().lt(10000) & values.ne(0),
        )
    else:
        unit_base = pd.Series(np.nan, index=df.index, dtype=float)

    quantity = (
        _quantity_or_one(df[quantity_col])
        if quantity_col is not None
        else pd.Series(1, index=df.index, dtype=int)
    )
    sku = (
        df[sku_col].astype("string").fillna("").str.strip()
        if sku_col is not None
        else pd.Series("", index=df.index, dtype="string")
    )
    reason = (
        df[reason_col].astype("string").fillna("").str.strip()
        if reason_col is not None
        else pd.Series("", index=df.index, dtype="string")
    )

    is_cancelled = status.str.upper().str.contains(
        r"BATAL|CANCEL", regex=True, na=False
    )
    lines = pd.DataFrame(
        {
            "_source_row": np.arange(len(df), dtype=int),
            "month": created_at.dt.to_period("M").dt.to_timestamp(),
            "order_date": created_at.dt.normalize(),
            "order_id": order_id,
            "sku": sku,
            "quantity": quantity,
            "subtotal": subtotal,
            "seller_rebate": seller_rebate,
            "shopee_discount": shopee_discount,
            "unit_base": unit_base,
            "order_status": status,
            "cancellation_reason": reason,
            "is_mfr_included": ~is_cancelled,
        }
    )
    lines = lines[
        lines["month"].notna()
        & lines["order_id"].ne("")
        & lines["is_mfr_included"]
    ].copy()
    if lines.empty:
        return _empty_order_detail()

    group_columns = ["month", "order_id"]
    group = lines.groupby(group_columns, sort=False)
    total_quantity = group["quantity"].transform("sum")
    order_rebate = group["seller_rebate"].transform("max")
    order_shopee_discount = group["shopee_discount"].transform("sum")
    lines["_order_seller_rebate"] = order_rebate
    lines["_order_shopee_discount"] = order_shopee_discount
    target_amount = (
        group["subtotal"].transform("sum")
        - order_rebate
        + order_shopee_discount
    )
    fallback_unit_base = lines["subtotal"] / lines["quantity"]
    reference_unit_base = lines["unit_base"].where(
        lines["unit_base"].notna() & lines["unit_base"].ne(0),
        fallback_unit_base,
    )
    exact_unit_amount = (
        reference_unit_base
        + (lines["shopee_discount"] / lines["quantity"])
        - (order_rebate / total_quantity)
    )
    expanded = _expand_and_reconcile_paid_amount(
        lines,
        group_columns=group_columns,
        exact_unit_amount=exact_unit_amount,
        target_amount=target_amount,
    )
    expanded["marketplace"] = "SHO"
    expanded["sequence"] = (
        expanded.groupby(group_columns, sort=False).cumcount() + 1
    ).astype(int)
    first_order_row = expanded["sequence"].eq(1)
    expanded["seller_rebate"] = np.where(
        first_order_row,
        expanded["_order_seller_rebate"],
        0.0,
    )
    expanded["shopee_discount"] = np.where(
        first_order_row,
        expanded["_order_shopee_discount"],
        0.0,
    )
    expanded["platform_discount"] = 0.0
    expanded["is_mfr_included"] = True
    return _standardize_order_detail_schema(
        expanded.reset_index(drop=True),
        "detail pesanan Shopee",
    )

def process_shopee_mfr(df: pd.DataFrame) -> dict[str, Any]:
    """
    Shopee MFR per month:
      Net GMV = Subtotal Pesanan (non-batal)
                - Voucher Ditanggung Penjual (order-level deduplicated)
                + Diskon Dari Shopee (order-level deduplicated)

    Eligibility:
      - Status Pesanan bukan BATAL / CANCEL / CANCELED / CANCELLED (case-insensitive)
      - Status aktif lainnya (Selesai, Sedang Dikirim, Telah Dikirim, Pesanan Diterima, dll.) = INCLUDED
    """
    required = {
        "No. Pesanan": (),
        "Status Pesanan": ("Order Status", "Status"),
        "Waktu Pesanan Dibuat": (),
        "Subtotal Pesanan": (),
        "Voucher Ditanggung Penjual": (),
        "Diskon Dari Shopee": (),
    }
    cols = _require_columns(df, "SHO", required)

    order_id = _normalized_id_series(df[cols["No. Pesanan"]])
    status_raw = df[cols["Status Pesanan"]].astype("string").fillna("").str.strip()
    status_upper = status_raw.str.upper()

    month = _month_start(df[cols["Waktu Pesanan Dibuat"]], dayfirst=False)
    
    # Shopee number normalization:
    raw_subtotal = df[cols["Subtotal Pesanan"]]
    raw_discount = df[cols["Diskon Dari Shopee"]]
    raw_voucher = df[cols["Voucher Ditanggung Penjual"]]

    # Apply the x1,000 rule only to genuinely numeric source cells. Text such as
    # "190.000" is already IDR-formatted and must become 190000 exactly once.
    subtotal = _normalize_shopee_amount(
        raw_subtotal,
        multiply_numeric=lambda values: values.abs().lt(10000) & values.ne(0),
    )

    voucher_seller = _to_idr_number(raw_voucher)

    shopee_discount = _normalize_shopee_amount(
        raw_discount,
        multiply_numeric=lambda values: values.gt(0) & values.lt(100),
    )

    is_cancel = status_upper.str.contains(r"BATAL|CANCEL", regex=True, na=False)
    not_cancel = ~is_cancel

    # Base work dataframe
    work = pd.DataFrame({
        "order_id": order_id,
        "status": status_raw,
        "month": month,
        "subtotal": subtotal,
        "voucher_seller": voucher_seller,
        "shopee_discount": shopee_discount,
        "is_cancel": is_cancel,
        "not_cancel": not_cancel,
    })

    work = work[work["month"].notna()].copy()
    if work.empty:
        return _empty_mfr_dict("SHO")

    # Order-level deduplication:
    # Subtotal Pesanan -> SUM
    # Voucher Ditanggung Penjual -> MAX
    # Diskon Dari Shopee -> SUM
    order_level = (
        work.groupby(["month", "order_id"], as_index=False)
        .agg(
            subtotal_all=("subtotal", "sum"),
            not_cancel=("not_cancel", "first"),
            is_cancel=("is_cancel", "first"),
            voucher_seller=("voucher_seller", "max"),
            shopee_discount=("shopee_discount", "sum"),
            status=("status", "first"),
            rows=("subtotal", "size"),
        )
    )
    order_level["gmv_ex_cancel"] = np.where(
        order_level["not_cancel"], order_level["subtotal_all"], 0.0
    )
    order_level["seller_rebate"] = np.where(
        order_level["not_cancel"], order_level["voucher_seller"], 0.0
    )
    order_level["shopee_discount_mfr"] = np.where(
        order_level["not_cancel"], order_level["shopee_discount"], 0.0
    )
    order_level["net_gmv"] = (
        order_level["gmv_ex_cancel"]
        - order_level["seller_rebate"]
        + order_level["shopee_discount_mfr"]
    )

    # Monthly aggregation
    monthly = (
        order_level.groupby("month", as_index=False)
        .agg(
            gmv_all=("subtotal_all", "sum"),
            gmv_ex_cancel=("gmv_ex_cancel", "sum"),
            seller_rebate=("seller_rebate", "sum"),
            shopee_discount=("shopee_discount_mfr", "sum"),
            net_gmv=("net_gmv", "sum"),
            source_rows=("rows", "sum"),
            included_rows=("not_cancel", lambda s: int(order_level.loc[s.index, "rows"][s].sum())),
            excluded_rows=("is_cancel", lambda s: int(order_level.loc[s.index, "rows"][s].sum())),
        )
    )
    monthly["platform_discount"] = 0.0
    monthly["marketplace"] = "SHO"

    monthly = monthly[MFR_MONTHLY_COLUMNS]

    # Status breakdown per month
    status_records = []
    for (m, st, is_st_cancel), grp in order_level.groupby(
        ["month", "status", "is_cancel"], dropna=False
    ):
        treatment = "Excluded" if bool(is_st_cancel) else "Included"
        status_records.append({
            "marketplace": "SHO",
            "month": m,
            "status": st if str(st).strip() != "" else "(Blank)",
            "rows": int(grp["rows"].sum()),
            "net_gmv": float(grp["net_gmv"].sum()),
            "treatment": treatment,
        })

    status_breakdown = pd.DataFrame(status_records)

    warnings: list[str] = []
    if (status_raw == "").any():
        blank_cnt = int((status_raw == "").sum())
        warnings.append(f"SHO: Terdapat {blank_cnt} baris dengan Status Pesanan kosong.")

    return {
        "marketplace": "SHO",
        "monthly": monthly,
        "status_breakdown": status_breakdown,
        "order_detail": build_shopee_order_detail(df),
        "warnings": warnings,
        "raw_df": df,
    }


# ============================================================
# 2. TIKTOK / TOKOPEDIA MFR
# ============================================================

def _tik_family_cancel_mask(
    sub_status: pd.Series,
    order_status: pd.Series,
    cancel_type: pd.Series,
) -> pd.Series:
    sub_status_up = sub_status.astype("string").fillna("").str.strip().str.upper()
    order_status_up = order_status.astype("string").fillna("").str.strip().str.upper()
    cancel_type_up = cancel_type.astype("string").fillna("").str.strip().str.upper()
    return (
        sub_status_up.eq("DIBATALKAN")
        | order_status_up.str.contains(
            r"BATAL|CANCEL|CANCELED|CANCELLED", regex=True, na=False
        )
        | cancel_type_up.str.contains("CANCEL", regex=False, na=False)
    )


def build_tik_family_order_detail(
    df: pd.DataFrame,
    marketplace: str,
) -> pd.DataFrame:
    """Build eligible TikTok/Tokopedia unit detail using Purchase Channel."""
    marketplace = marketplace.upper()
    if marketplace not in {"TIK", "TOK"}:
        raise BenchmarkInputError("Marketplace TikTok-family harus TIK atau TOK.")

    required = {
        "Order ID": (),
        "Order Sub Status": ("Order Substatus", "Order Sub-Status"),
        "Order Status": (),
        "Cancelation/Return Type": ("Cancellation/Return Type", "Cancel Type"),
        "Created Time": (),
        "SKU Subtotal After Discount": (),
        "SKU Platform Discount": (),
    }
    cols = _require_columns(df, marketplace, required)
    channel_col = _resolve_column(df, "Purchase Channel", ["Channel"])
    sku_col = _resolve_column(
        df,
        "Seller SKU",
        ["SKU Penjual", "Merchant SKU", "SKU"],
    )
    quantity_col = _resolve_column(df, "Quantity", ["Qty", "Jumlah"])
    reason_col = _resolve_column(
        df,
        "Cancel Reason",
        ["Cancellation Reason", "Cancelation Reason", "Alasan Pembatalan"],
    )

    order_id = _normalized_id_series(df[cols["Order ID"]])
    sub_status = df[cols["Order Sub Status"]].astype("string").fillna("").str.strip()
    order_status = df[cols["Order Status"]].astype("string").fillna("").str.strip()
    cancel_type = (
        df[cols["Cancelation/Return Type"]]
        .astype("string")
        .fillna("")
        .str.strip()
    )
    created_at = _parse_datetime(df[cols["Created Time"]], dayfirst=True)
    subtotal = _to_idr_number(df[cols["SKU Subtotal After Discount"]])
    platform_discount = _to_idr_number(df[cols["SKU Platform Discount"]])
    purchase_channel = (
        df[channel_col].astype("string").fillna("").str.strip().str.upper()
        if channel_col is not None
        else pd.Series(
            "TIKTOK" if marketplace == "TIK" else "TOKOPEDIA",
            index=df.index,
            dtype="string",
        )
    )
    target_channel = "TIKTOK" if marketplace == "TIK" else "TOKOPEDIA"
    quantity = (
        _quantity_or_one(df[quantity_col])
        if quantity_col is not None
        else pd.Series(1, index=df.index, dtype=int)
    )
    sku = (
        df[sku_col].astype("string").fillna("").str.strip()
        if sku_col is not None
        else pd.Series("", index=df.index, dtype="string")
    )
    reason = (
        df[reason_col].astype("string").fillna("").str.strip()
        if reason_col is not None
        else pd.Series("", index=df.index, dtype="string")
    )
    is_cancelled = _tik_family_cancel_mask(sub_status, order_status, cancel_type)

    lines = pd.DataFrame(
        {
            "_source_row": np.arange(len(df), dtype=int),
            "month": created_at.dt.to_period("M").dt.to_timestamp(),
            "order_date": created_at.dt.normalize(),
            "order_id": order_id,
            "sku": sku,
            "quantity": quantity,
            "line_total": subtotal + platform_discount,
            "platform_discount": platform_discount,
            "purchase_channel": purchase_channel,
            "order_status": order_status,
            "cancellation_reason": reason,
            "is_mfr_included": ~is_cancelled,
        }
    )
    lines = lines[
        lines["month"].notna()
        & lines["order_id"].ne("")
        & lines["purchase_channel"].eq(target_channel)
        & lines["is_mfr_included"]
    ].copy()
    if lines.empty:
        return _empty_order_detail()

    # Reference: day ascending, with the source export order retained inside a day.
    lines = lines.sort_values("order_date", kind="stable").reset_index(drop=True)
    group_columns = ["month", "order_id", "purchase_channel"]
    group = lines.groupby(group_columns, sort=False)
    target_amount = group["line_total"].transform("sum")
    exact_unit_amount = lines["line_total"] / lines["quantity"]
    expanded = _expand_and_reconcile_paid_amount(
        lines,
        group_columns=group_columns,
        exact_unit_amount=exact_unit_amount,
        target_amount=target_amount,
    )
    expanded["marketplace"] = expanded["purchase_channel"].map(
        {"TIKTOK": "TIK", "TOKOPEDIA": "TOK"}
    )
    expanded["sequence"] = (
        expanded.groupby(group_columns, sort=False).cumcount() + 1
    ).astype(int)
    expanded["seller_rebate"] = 0.0
    expanded["shopee_discount"] = 0.0
    # SKU Platform Discount belongs to the source line. Quantity expansion is
    # retained for Harga Terbayarkan, so record the line value on its first
    # expanded unit only to prevent duplication while preserving the line sum.
    expanded["platform_discount"] = np.where(
        expanded["_unit_position"].eq(0),
        expanded["platform_discount"],
        0.0,
    )
    expanded["is_mfr_included"] = True
    return _standardize_order_detail_schema(
        expanded.reset_index(drop=True),
        f"detail pesanan {marketplace}",
    )

def process_tik_family_mfr(
    df: pd.DataFrame,
    marketplace: str,
) -> dict[str, Any]:
    """
    TikTok / Tokopedia MFR per month:
      Net GMV = GMV Exc Batal (SKU Subtotal After Discount)
                + Platform Discount (SKU Platform Discount)

    Purchase Channel Aware:
      - Jika kolom Purchase Channel tersedia:
        - TIK -> hanya Purchase Channel == 'TIKTOK'
        - TOK -> hanya Purchase Channel == 'TOKOPEDIA'

    Cancel Filter 3 Layer:
      Order EXCLUDED jika salah satu terpenuhi:
        1. Order Sub Status == 'DIBATALKAN' (atau Order Substatus)
        2. Order Status mengandung 'BATAL|CANCEL|CANCELED|CANCELLED'
        3. Cancelation/Return Type mengandung 'CANCEL'
      Status aktif lainnya (Selesai, Terkirim, Sedang transit, dll.) = INCLUDED

    Definisi GMV ALL:
      SUM(SKU Subtotal After Discount) + SUM(SKU Platform Discount)
      untuk SEMUA pesanan channel tersebut (termasuk order batal).
    """
    if marketplace not in {"TIK", "TOK"}:
        raise BenchmarkInputError("Marketplace TikTok-family harus TIK atau TOK.")

    required = {
        "Order ID": (),
        "Order Sub Status": ("Order Substatus", "Order Sub-Status"),
        "Order Status": (),
        "Cancelation/Return Type": ("Cancellation/Return Type", "Cancel Type"),
        "Created Time": (),
        "SKU Subtotal After Discount": (),
        "SKU Platform Discount": (),
    }
    cols = _require_columns(df, marketplace, required)

    # Resolve optional Purchase Channel column
    channel_col = _resolve_column(df, "Purchase Channel", ["Channel", "purchase channel"])

    order_id = _normalized_id_series(df[cols["Order ID"]])
    sub_status_raw = df[cols["Order Sub Status"]].astype("string").fillna("").str.strip()
    order_status_raw = df[cols["Order Status"]].astype("string").fillna("").str.strip()
    cancel_type_raw = df[cols["Cancelation/Return Type"]].astype("string").fillna("").str.strip()

    created_time = _parse_datetime(df[cols["Created Time"]], dayfirst=True)
    month = created_time.dt.to_period("M").dt.to_timestamp()

    subtotal = _to_idr_number(df[cols["SKU Subtotal After Discount"]])
    platform_discount = _to_idr_number(df[cols["SKU Platform Discount"]])

    # 1. Purchase Channel filtering
    if channel_col is not None:
        purchase_channel = df[channel_col].astype("string").fillna("").str.strip().str.upper()
        if marketplace == "TIK":
            channel_mask = purchase_channel.eq("TIKTOK")
        else:
            channel_mask = purchase_channel.eq("TOKOPEDIA")
    else:
        purchase_channel = pd.Series(
            "TIKTOK" if marketplace == "TIK" else "TOKOPEDIA",
            index=df.index,
            dtype="string",
        )
        channel_mask = pd.Series(True, index=df.index)

    # 2. 3-Layer Cancel Detection
    is_cancelled = _tik_family_cancel_mask(
        sub_status_raw,
        order_status_raw,
        cancel_type_raw,
    )

    is_valid_order_id = order_id.ne("")

    # The master contract does not require a non-blank substatus. Any order with
    # a valid ID is included unless one of the three cancellation layers matches.
    is_included = is_valid_order_id & (~is_cancelled) & channel_mask
    is_excluded = is_valid_order_id & is_cancelled & channel_mask

    work = pd.DataFrame({
        "order_id": order_id,
        "sub_status": sub_status_raw,
        "order_status": order_status_raw,
        "cancel_type": cancel_type_raw,
        "created_time": created_time,
        "month": month,
        "subtotal": subtotal,
        "platform_discount": platform_discount,
        "purchase_channel": purchase_channel,
        "channel_mask": channel_mask,
        "is_cancelled": is_cancelled,
        "is_included": is_included,
        "is_excluded": is_excluded,
    })

    # GMV ALL includes every row for the relevant channel, including canceled
    # orders. A valid Order ID is only required for eligible/order-level detail.
    channel_work = work[work["month"].notna() & work["channel_mask"]].copy()
    if channel_work.empty:
        return _empty_mfr_dict(marketplace)

    valid_work = channel_work[channel_work["order_id"].ne("")].copy()

    # An Order ID must resolve to one channel/date/status. Cancellation type can
    # legitimately vary by SKU, so cancellation itself is aggregated with ANY.
    consistency_columns = [
        "purchase_channel",
        "created_time",
        "order_status",
        "sub_status",
    ]
    if not valid_work.empty:
        grouped_orders = valid_work.groupby(
            ["order_id", "purchase_channel"], dropna=False
        )
        inconsistent_orders: set[str] = set()
        for column in consistency_columns:
            inconsistent = grouped_orders[column].nunique(dropna=False).gt(1)
            inconsistent_orders.update(
                str(index[0]) for index in inconsistent[inconsistent].index.tolist()
            )
        if inconsistent_orders:
            samples = sorted(inconsistent_orders)[:5]
            raise BenchmarkInputError(
                f"{marketplace}: status/channel/tanggal tidak konsisten untuk "
                f"{len(inconsistent_orders)} Order ID. Contoh: {samples}"
            )

    def first_nonblank(values: pd.Series) -> str:
        normalized = values.astype("string").fillna("").str.strip()
        nonblank = normalized[normalized.ne("")]
        return str(nonblank.iloc[0]) if not nonblank.empty else ""

    order_level = (
        valid_work.groupby(["order_id", "purchase_channel"], as_index=False)
        .agg(
            month=("month", "first"),
            sub_status=("sub_status", first_nonblank),
            order_status=("order_status", first_nonblank),
            cancel_type=("cancel_type", first_nonblank),
            subtotal=("subtotal", "sum"),
            platform_discount=("platform_discount", "sum"),
            is_cancelled=("is_cancelled", "max"),
            raw_rows=("order_id", "size"),
        )
    )
    order_level["is_included"] = ~order_level["is_cancelled"].astype(bool)
    order_level["is_excluded"] = order_level["is_cancelled"].astype(bool)
    order_level["status"] = order_level["sub_status"].where(
        order_level["sub_status"].ne(""), order_level["order_status"]
    )
    order_level["status"] = order_level["status"].replace("", "(Blank)")
    order_level["gmv_ex_cancel"] = np.where(
        order_level["is_included"], order_level["subtotal"], 0.0
    )
    order_level["platform_discount_mfr"] = np.where(
        order_level["is_included"], order_level["platform_discount"], 0.0
    )
    order_level["net_gmv"] = (
        order_level["gmv_ex_cancel"] + order_level["platform_discount_mfr"]
    )

    gmv_all = (
        channel_work.assign(
            row_gross=channel_work["subtotal"] + channel_work["platform_discount"]
        )
        .groupby("month", as_index=False)
        .agg(gmv_all=("row_gross", "sum"))
    )
    order_monthly = (
        order_level.groupby("month", as_index=False)
        .agg(
            gmv_ex_cancel=("gmv_ex_cancel", "sum"),
            platform_discount=("platform_discount_mfr", "sum"),
            net_gmv=("net_gmv", "sum"),
            source_rows=("order_id", "size"),
            included_rows=("is_included", "sum"),
            excluded_rows=("is_excluded", "sum"),
        )
    )
    monthly = gmv_all.merge(order_monthly, on="month", how="left")
    for column in [
        "gmv_ex_cancel",
        "platform_discount",
        "net_gmv",
        "source_rows",
        "included_rows",
        "excluded_rows",
    ]:
        monthly[column] = monthly[column].fillna(0)
    monthly["seller_rebate"] = 0.0
    monthly["shopee_discount"] = 0.0
    monthly["marketplace"] = marketplace

    monthly = monthly[MFR_MONTHLY_COLUMNS]

    # Status breakdown uses the same one-row-per-order detail as the calculation.
    status_records = []
    for (m, st, included), grp in order_level.groupby(
        ["month", "status", "is_included"], dropna=False
    ):
        treatment = "Included" if bool(included) else "Excluded"
        st_rows = len(grp)
        status_records.append({
            "marketplace": marketplace,
            "month": m,
            "status": st if str(st).strip() != "" else "(Blank)",
            "rows": st_rows,
            "net_gmv": float(grp["net_gmv"].sum()),
            "treatment": treatment,
        })

    status_breakdown = pd.DataFrame(status_records)

    warnings: list[str] = []
    blank_cnt = int(channel_work["sub_status"].eq("").sum())
    if blank_cnt > 0:
        warnings.append(
            f"{marketplace}: Terdapat {blank_cnt} baris dengan Order Sub Status kosong "
            "(tetap diproses berdasarkan full cancel rule)."
        )
    blank_id_cnt = int(channel_work["order_id"].eq("").sum())
    if blank_id_cnt > 0:
        warnings.append(
            f"{marketplace}: Terdapat {blank_id_cnt} baris dengan Order ID kosong; "
            "nilai hanya masuk GMV ALL dan tidak masuk Net GMV."
        )

    return {
        "marketplace": marketplace,
        "monthly": monthly,
        "status_breakdown": status_breakdown,
        "order_detail": build_tik_family_order_detail(df, marketplace),
        "warnings": warnings,
        "raw_df": df,
    }


# ============================================================
# 3. LAZADA MFR
# ============================================================

def build_lazada_order_detail(df: pd.DataFrame) -> pd.DataFrame:
    """Build eligible Lazada line detail; no quantity expansion without a sample."""
    required = {
        "orderItemId": (),
        "createTime": (),
        "unitPrice": (),
        "platformDiscountTotal": (),
        "status": (),
    }
    cols = _require_columns(df, "LAZ", required)
    order_col = _resolve_column(
        df,
        "orderId",
        ["orderNumber", "Order Number", "Order ID", "order_id"],
    ) or cols["orderItemId"]
    sku_col = _resolve_column(
        df,
        "sellerSku",
        ["Seller SKU", "SKU Penjual", "SKU"],
    )
    reason_col = _resolve_column(
        df,
        "cancelReason",
        ["Cancel Reason", "Cancellation Reason", "failureReason"],
    )

    order_id = _normalized_id_series(df[order_col])
    created_at = _parse_datetime(df[cols["createTime"]], dayfirst=True)
    status = df[cols["status"]].astype("string").fillna("").str.strip()
    unit_price = _to_idr_number(df[cols["unitPrice"]])
    platform_discount = _to_idr_number(df[cols["platformDiscountTotal"]])
    sku = (
        df[sku_col].astype("string").fillna("").str.strip()
        if sku_col is not None
        else pd.Series("", index=df.index, dtype="string")
    )
    reason = (
        df[reason_col].astype("string").fillna("").str.strip()
        if reason_col is not None
        else pd.Series("", index=df.index, dtype="string")
    )
    is_included = ~status.str.lower().isin({"canceled", "cancelled"})
    detail = pd.DataFrame(
        {
            "order_date": created_at.dt.normalize(),
            "marketplace": "LAZ",
            "order_id": order_id,
            "sku": sku,
            "paid_amount": unit_price + platform_discount,
            "seller_rebate": 0.0,
            "shopee_discount": 0.0,
            "platform_discount": platform_discount,
            "order_status": status,
            "cancellation_reason": reason,
            "is_mfr_included": is_included,
        }
    )
    detail = detail[
        detail["order_date"].notna()
        & detail["order_id"].ne("")
        & detail["is_mfr_included"]
    ].copy()
    if detail.empty:
        return _empty_order_detail()
    detail["sequence"] = (
        detail.groupby("order_id", sort=False).cumcount() + 1
    ).astype(int)
    return _standardize_order_detail_schema(
        detail.reset_index(drop=True),
        "detail pesanan Lazada",
    )

def process_lazada_mfr(df: pd.DataFrame) -> dict[str, Any]:
    """
    Lazada MFR per month:
      Net GMV = unitPrice (non-cancel) + platformDiscountTotal (non-cancel)

    Eligibility:
      - status bukan canceled / cancelled (case-insensitive)
      - status aktif lainnya = INCLUDED
    """
    required = {
        "orderItemId": (),
        "createTime": (),
        "unitPrice": (),
        "platformDiscountTotal": (),
        "status": (),
    }
    cols = _require_columns(df, "LAZ", required)

    status_raw = df[cols["status"]].astype("string").fillna("").str.strip()
    status_lower = status_raw.str.lower()

    created_time = _parse_datetime(df[cols["createTime"]], dayfirst=True)
    month = created_time.dt.to_period("M").dt.to_timestamp()

    unit_price = _to_idr_number(df[cols["unitPrice"]])
    platform_discount = _to_idr_number(df[cols["platformDiscountTotal"]])

    is_cancel = status_lower.isin({"canceled", "cancelled"})
    not_cancel = ~is_cancel

    work = pd.DataFrame({
        "status": status_raw,
        "month": month,
        "unit_price": unit_price,
        "platform_discount": platform_discount,
        "is_cancel": is_cancel,
        "not_cancel": not_cancel,
    })

    work = work[work["month"].notna()].copy()
    if work.empty:
        return _empty_mfr_dict("LAZ")

    work["row_gross"] = work["unit_price"] + work["platform_discount"]
    work["gmv_ex_cancel"] = np.where(
        work["not_cancel"], work["unit_price"], 0.0
    )
    work["discount_mfr"] = np.where(
        work["not_cancel"], work["platform_discount"], 0.0
    )
    work["net_gmv"] = np.where(work["not_cancel"], work["row_gross"], 0.0)

    # Monthly aggregation
    monthly = (
        work.groupby("month", as_index=False)
        .agg(
            gmv_all=("row_gross", "sum"),
            gmv_ex_cancel=("gmv_ex_cancel", "sum"),
            platform_discount=("discount_mfr", "sum"),
            net_gmv=("net_gmv", "sum"),
            source_rows=("unit_price", "size"),
            included_rows=("not_cancel", "sum"),
            excluded_rows=("is_cancel", "sum"),
        )
    )
    monthly["seller_rebate"] = 0.0
    monthly["shopee_discount"] = 0.0
    monthly["marketplace"] = "LAZ"

    monthly = monthly[MFR_MONTHLY_COLUMNS]

    # Status breakdown per month
    status_records = []
    for (m, st), grp in work.groupby(["month", "status"]):
        st_lower = str(st).lower().strip()
        treatment = "Excluded" if st_lower in {"canceled", "cancelled"} else "Included"
        status_records.append({
            "marketplace": "LAZ",
            "month": m,
            "status": st if st != "" else "(Blank)",
            "rows": len(grp),
            "net_gmv": float(grp["net_gmv"].sum()),
            "treatment": treatment,
        })

    status_breakdown = pd.DataFrame(status_records)

    warnings: list[str] = []
    if (status_raw == "").any():
        blank_cnt = int((status_raw == "").sum())
        warnings.append(f"LAZ: Terdapat {blank_cnt} baris dengan status kosong.")

    return {
        "marketplace": "LAZ",
        "monthly": monthly,
        "status_breakdown": status_breakdown,
        "order_detail": build_lazada_order_detail(df),
        "warnings": warnings,
        "raw_df": df,
    }


def _empty_mfr_dict(marketplace: str) -> dict[str, Any]:
    empty_monthly = pd.DataFrame(columns=MFR_MONTHLY_COLUMNS)
    empty_status = pd.DataFrame(columns=MFR_STATUS_COLUMNS)
    return {
        "marketplace": marketplace,
        "monthly": empty_monthly,
        "status_breakdown": empty_status,
        "order_detail": _empty_order_detail(),
        "warnings": [],
        "raw_df": pd.DataFrame(),
    }


# ============================================================
# 4. DISPATCHER & CALCULATION
# ============================================================

def process_dataframe_mfr(
    df: pd.DataFrame,
    marketplace: str,
    filename: str | None = None,
) -> dict[str, Any]:
    marketplace = marketplace.upper()
    if df is None or df.empty:
        return _empty_mfr_dict(marketplace)

    if marketplace == "SHO":
        return process_shopee_mfr(df)
    if marketplace in {"TIK", "TOK"}:
        return process_tik_family_mfr(df, marketplace)
    if marketplace == "LAZ":
        return process_lazada_mfr(df)

    raise BenchmarkInputError(f"Marketplace '{marketplace}' belum didukung untuk MFR.")


def calculate_mfr(
    file_results: Iterable[dict[str, Any]],
    selected_month: pd.Timestamp | str,
) -> dict[str, Any]:
    """
    Kalkulasi MFR Final untuk 1 Bulan Terpilih:
      Total Net GMV = SHO Net GMV + TIK Net GMV + TOK Net GMV + LAZ Net GMV
      TIDAK ADA AVERAGE.

    Returns:
      dict with:
        - month: pd.Timestamp
        - month_label: e.g. "August 2026"
        - mfr_total: float
        - total_rows: int (included rows)
        - marketplace_count: int
        - marketplace_summary: pd.DataFrame
        - marketplace_detail: pd.DataFrame
        - status_breakdown: pd.DataFrame
        - warnings: list[str]
        - audit: list[dict]
    """
    ts_month = pd.Timestamp(selected_month).to_period("M").to_timestamp()
    month_label = ts_month.strftime("%B %Y")

    all_monthly: list[pd.DataFrame] = []
    all_status: list[pd.DataFrame] = []
    all_order_detail: list[pd.DataFrame] = []
    all_warnings: list[str] = []

    for res in file_results:
        if not res:
            continue
        m_df = res.get("monthly")
        s_df = res.get("status_breakdown")
        d_df = res.get("order_detail")
        w_list = res.get("warnings", [])

        if m_df is not None and not m_df.empty:
            _require_internal_schema(m_df, MFR_MONTHLY_COLUMNS, "hasil file bulanan")
            all_monthly.append(m_df)
        if s_df is not None and not s_df.empty:
            _require_internal_schema(s_df, MFR_STATUS_COLUMNS, "breakdown status file")
            all_status.append(s_df)
        if d_df is not None and not d_df.empty:
            all_order_detail.append(
                _standardize_order_detail_schema(
                    d_df,
                    f"daftar pesanan file {res.get('marketplace', '')}",
                )
            )
        all_warnings.extend(w_list)

    if not all_monthly:
        raise BenchmarkInputError("Tidak ada data bulanan MFR yang valid.")

    combined_monthly = pd.concat(all_monthly, ignore_index=True)
    combined_monthly["month"] = pd.to_datetime(combined_monthly["month"])
    numeric_columns = [
        *MFR_METRIC_COLUMNS,
        "source_rows",
        "included_rows",
        "excluded_rows",
    ]
    for column in numeric_columns:
        try:
            combined_monthly[column] = pd.to_numeric(
                combined_monthly[column], errors="raise"
            )
        except (TypeError, ValueError) as exc:
            raise BenchmarkInputError(
                f"Schema MFR tidak valid: kolom internal '{column}' harus numeric."
            ) from exc

    # Filter strictly for selected month
    month_filter = combined_monthly["month"] == ts_month
    filtered_monthly = combined_monthly[month_filter].copy()

    if filtered_monthly.empty:
        raise BenchmarkInputError(
            f"Tidak ada data transaksi pada bulan yang dipilih: {month_label}."
        )

    # Group by marketplace for selected month
    marketplace_detail = (
        filtered_monthly.groupby("marketplace", as_index=False)
        .agg(
            gmv_all=("gmv_all", "sum"),
            gmv_ex_cancel=("gmv_ex_cancel", "sum"),
            seller_rebate=("seller_rebate", "sum"),
            shopee_discount=("shopee_discount", "sum"),
            platform_discount=("platform_discount", "sum"),
            net_gmv=("net_gmv", "sum"),
            source_rows=("source_rows", "sum"),
            included_rows=("included_rows", "sum"),
            excluded_rows=("excluded_rows", "sum"),
        )
        .sort_values("marketplace")
        .reset_index(drop=True)
    )

    mfr_total = float(marketplace_detail["net_gmv"].sum())
    total_included_rows = int(marketplace_detail["included_rows"].sum())
    marketplace_count = int(marketplace_detail["marketplace"].nunique())

    # Keep calculation results canonical. Display labels are applied only by UI
    # and Excel-report copies.
    marketplace_summary = marketplace_detail[
        ["marketplace", *MFR_METRIC_COLUMNS]
    ].copy()
    total_row: dict[str, Any] = {"marketplace": "TOTAL MFR"}
    total_row.update(
        {
            column: float(marketplace_detail[column].sum())
            for column in MFR_METRIC_COLUMNS
        }
    )
    marketplace_summary = pd.concat(
        [marketplace_summary, pd.DataFrame([total_row])], ignore_index=True
    )

    # Consolidated Status Breakdown for selected month
    if all_status:
        combined_status = pd.concat(all_status, ignore_index=True)
        combined_status["month"] = pd.to_datetime(combined_status["month"])
        status_filtered = combined_status[combined_status["month"] == ts_month].copy()

        status_breakdown = (
            status_filtered.groupby(["marketplace", "status", "treatment"], as_index=False)
            .agg(
                rows=("rows", "sum"),
                net_gmv=("net_gmv", "sum"),
            )
            .sort_values(
                by=["marketplace", "treatment", "net_gmv"],
                ascending=[True, True, False],
            )
            .reset_index(drop=True)
        )
    else:
        status_breakdown = pd.DataFrame(
            columns=["marketplace", "status", "treatment", "rows", "net_gmv"]
        )

    if all_order_detail:
        combined_detail = _standardize_order_detail_schema(
            pd.concat(all_order_detail, ignore_index=True),
            "hasil concat daftar pesanan",
        )
        combined_detail["order_date"] = pd.to_datetime(
            combined_detail["order_date"], errors="coerce"
        )
        order_detail = combined_detail[
            combined_detail["order_date"].dt.to_period("M").dt.to_timestamp()
            == ts_month
        ].copy()
        marketplace_rank = {"SHO": 0, "TIK": 1, "TOK": 2, "LAZ": 3}
        order_detail["_marketplace_rank"] = (
            order_detail["marketplace"].map(marketplace_rank).fillna(99)
        )
        order_detail = (
            order_detail.sort_values(
                ["_marketplace_rank", "order_date"], kind="stable"
            )
            .drop(columns="_marketplace_rank")
            .reset_index(drop=True)
        )
        order_detail = _standardize_order_detail_schema(
            order_detail,
            "output daftar pesanan",
        )
    else:
        order_detail = _empty_order_detail()

    # Validate MFR results
    audit = validate_mfr_calculation(marketplace_detail, status_breakdown, mfr_total)
    audit.extend(validate_mfr_order_detail(marketplace_detail, order_detail))

    return {
        "month": ts_month,
        "month_label": month_label,
        "mfr_total": mfr_total,
        "total_rows": total_included_rows,
        "marketplace_count": marketplace_count,
        "marketplace_summary": marketplace_summary,
        "marketplace_detail": marketplace_detail,
        "status_breakdown": status_breakdown,
        "order_detail": order_detail,
        "warnings": list(dict.fromkeys(all_warnings)),
        "audit": audit,
    }


# ============================================================
# 5. ZERO MISTAKE AUDIT FOR MFR
# ============================================================

def validate_mfr_file(
    df: pd.DataFrame,
    marketplace: str,
    file_result: dict[str, Any],
) -> list[dict[str, Any]]:
    """
    Audit Zero Mistake untuk satu file MFR uploaded.
    """
    marketplace = marketplace.upper()
    audit: list[dict[str, Any]] = []
    tol = 0.0

    def add(status: str, check: str, detail: str, actual: Any = "", expected: Any = ""):
        audit.append(_audit_row(status, check, detail, actual=actual, expected=expected))

    monthly = file_result.get("monthly")
    if monthly is None or monthly.empty:
        add("FAIL", "Monthly MFR tersedia", "Tidak ada baris bulanan yang berhasil dihitung.")
        return audit

    # 1. Parsing tanggal
    if marketplace == "SHO":
        date_col = _resolve_column(df, "Waktu Pesanan Dibuat")
    elif marketplace in {"TIK", "TOK"}:
        date_col = _resolve_column(df, "Created Time")
    elif marketplace == "LAZ":
        date_col = _resolve_column(df, "createTime")
    else:
        date_col = None

    if date_col:
        raw_dates = df[date_col].astype("string").fillna("").str.strip()
        nonempty_date = raw_dates.ne("")
        dayfirst = marketplace != "SHO"
        parsed_dates = _parse_datetime(df[date_col], dayfirst=dayfirst)
        invalid_dates = int((nonempty_date & parsed_dates.isna()).sum())
        add(
            "PASS" if invalid_dates == 0 else "FAIL",
            f"Parsing tanggal {marketplace}",
            f"Semua nilai {date_col} non-kosong harus dapat diparse menjadi tanggal.",
            actual=invalid_dates,
            expected=0,
        )

    # 2. Parsing nominal
    numeric_fields: list[str] = []
    if marketplace == "SHO":
        numeric_fields = ["Subtotal Pesanan", "Voucher Ditanggung Penjual", "Diskon Dari Shopee"]
    elif marketplace in {"TIK", "TOK"}:
        numeric_fields = ["SKU Subtotal After Discount", "SKU Platform Discount"]
    elif marketplace == "LAZ":
        numeric_fields = ["unitPrice", "platformDiscountTotal"]

    for nf in numeric_fields:
        c = _resolve_column(df, nf)
        if c:
            invalid_num, samples = _nonempty_invalid_numeric_count(df[c])
            add(
                "PASS" if invalid_num == 0 else "FAIL",
                f"Parsing nominal {nf}",
                "Nominal tidak boleh diam-diam bernilai 0 karena format gagal dikenali."
                + (f" Contoh: {samples}" if samples else ""),
                actual=invalid_num,
                expected=0,
            )

    # 3. Order ID valid
    id_col = None
    if marketplace == "SHO":
        id_col = _resolve_column(df, "No. Pesanan")
    elif marketplace in {"TIK", "TOK"}:
        id_col = _resolve_column(df, "Order ID")
    elif marketplace == "LAZ":
        id_col = _resolve_column(df, "orderItemId")

    if id_col:
        id_series = _normalized_id_series(df[id_col])
        blank_ids = int(id_series.eq("").sum())
        add(
            "PASS" if blank_ids == 0 else "WARNING",
            f"Order ID valid {marketplace}",
            f"Semua baris diharapkan memiliki {id_col} non-kosong.",
            actual=blank_ids,
            expected=0,
        )

    # 4. Status blank warning
    st_col = None
    if marketplace == "SHO":
        st_col = _resolve_column(df, "Status Pesanan")
    elif marketplace in {"TIK", "TOK"}:
        st_col = _resolve_column(df, "Order Sub Status", ["Order Substatus"])
    elif marketplace == "LAZ":
        st_col = _resolve_column(df, "status")

    if st_col:
        raw_st = df[st_col].astype("string").fillna("").str.strip()
        blank_st = int(raw_st.eq("").sum())
        add(
            "PASS" if blank_st == 0 else "WARNING",
            f"Status pesanan terisi {marketplace}",
            f"Status pesanan pada kolom {st_col} tidak boleh kosong.",
            actual=blank_st,
            expected=0,
        )

    # 5. Formula check per row in monthly
    for _, row in monthly.iterrows():
        mp = row["marketplace"]
        if mp == "SHO":
            expected_omzet = (
                row["gmv_ex_cancel"]
                - row["seller_rebate"]
                + row["shopee_discount"]
            )
        else:
            expected_omzet = row["gmv_ex_cancel"] + row["platform_discount"]

        diff = abs(float(row["net_gmv"]) - float(expected_omzet))
        if diff > tol:
            add(
                "FAIL",
                f"Formula Net GMV {mp}",
                "Rekonsiliasi formula Net GMV tidak sesuai.",
                actual=f"Rp {float(row['net_gmv']):,.0f}",
                expected=f"Rp {float(expected_omzet):,.0f}",
            )
            break
    else:
        add(
            "PASS",
            f"Formula Net GMV {marketplace}",
            "Net GMV bulanan sesuai formula masing-masing marketplace.",
            actual="Sesuai",
            expected="Sesuai",
        )

    return audit


def validate_mfr_calculation(
    marketplace_detail: pd.DataFrame,
    status_breakdown: pd.DataFrame,
    mfr_total: float,
) -> list[dict[str, Any]]:
    """
    Audit Zero Mistake konsolidasi MFR untuk 1 bulan terpilih.
    """
    audit: list[dict[str, Any]] = []
    tol = 0.0

    def add(status: str, check: str, detail: str, actual: Any = "", expected: Any = ""):
        audit.append(_audit_row(status, check, detail, actual=actual, expected=expected))

    # 1. Total Net GMV = SUM(marketplace Net GMV)
    _require_internal_schema(
        marketplace_detail,
        ["marketplace", *MFR_METRIC_COLUMNS],
        "detail marketplace final",
    )
    sum_marketplace_mfr = float(marketplace_detail["net_gmv"].sum())
    total_diff = abs(mfr_total - sum_marketplace_mfr)
    add(
        "PASS" if total_diff <= tol else "FAIL",
        "Total Net GMV = Sum Marketplace",
        "Total Net GMV harus tepat sama dengan penjumlahan Net GMV seluruh marketplace tanpa average.",
        actual=f"Rp {mfr_total:,.0f}",
        expected=f"Rp {sum_marketplace_mfr:,.0f}",
    )

    # 2. Formula tiap marketplace
    for _, row in marketplace_detail.iterrows():
        mp = row["marketplace"]
        if mp == "SHO":
            expected = (
                row["gmv_ex_cancel"]
                - row["seller_rebate"]
                + row["shopee_discount"]
            )
        else:
            expected = row["gmv_ex_cancel"] + row["platform_discount"]
        diff = abs(float(row["net_gmv"]) - float(expected))
        add(
            "PASS" if diff <= tol else "FAIL",
            f"Formula Net GMV {mp}",
            f"Rekonsiliasi komponen detail {mp} terhadap Net GMV.",
            actual=f"Rp {float(row['net_gmv']):,.0f}",
            expected=f"Rp {float(expected):,.0f}",
        )

    # 3. Status breakdown reconciliation
    if not status_breakdown.empty:
        _require_internal_schema(
            status_breakdown,
            ["marketplace", "status", "treatment", "rows", "net_gmv"],
            "breakdown status final",
        )
        included_status = status_breakdown[
            status_breakdown["treatment"] == "Included"
        ]
        sum_included_omzet = float(included_status["net_gmv"].sum())
        st_diff = abs(mfr_total - sum_included_omzet)
        add(
            "PASS" if st_diff <= tol else "WARNING",
            "Reconcile Breakdown Status",
            "SUM Omzet dari seluruh status dengan treatment 'Included' harus merefleksikan Total Net GMV.",
            actual=f"Rp {sum_included_omzet:,.0f}",
            expected=f"Rp {mfr_total:,.0f}",
        )

        # Check canceled rows are strictly Excluded
        cancel_leak = status_breakdown[
            status_breakdown["status"]
            .astype(str)
            .str.upper()
            .str.contains(r"BATAL|CANCEL", regex=True, na=False)
            & (status_breakdown["treatment"] == "Included")
        ]
        add(
            "PASS" if cancel_leak.empty else "FAIL",
            "Canceled Orders Excluded",
            "Tidak boleh ada status pembatalan yang berstatus 'Included'.",
            actual=f"{len(cancel_leak)} status bocor" if not cancel_leak.empty else "0",
            expected="0",
        )

    return audit


def validate_mfr_order_detail(
    marketplace_detail: pd.DataFrame,
    order_detail: pd.DataFrame,
) -> list[dict[str, Any]]:
    """Reconcile eligible Daftar Pesanan rows to marketplace Net GMV exactly."""
    audit: list[dict[str, Any]] = []
    _require_internal_schema(
        marketplace_detail,
        ["marketplace", "net_gmv"],
        "rekonsiliasi daftar pesanan",
    )
    if order_detail.empty:
        for _, row in marketplace_detail.iterrows():
            audit.append(
                _audit_row(
                    "FAIL",
                    f"Daftar Pesanan = Net GMV {row['marketplace']}",
                    "Detail pesanan tidak tersedia untuk marketplace dengan Net GMV.",
                    actual="Rp 0",
                    expected=f"Rp {float(row['net_gmv']):,.0f}",
                )
            )
        return audit

    _require_internal_schema(
        order_detail,
        MFR_ORDER_DETAIL_COLUMNS,
        "rekonsiliasi daftar pesanan",
    )
    eligible = order_detail[order_detail["is_mfr_included"].astype(bool)]
    detail_metrics = [
        "paid_amount",
        "seller_rebate",
        "shopee_discount",
        "platform_discount",
    ]
    totals_by_marketplace = eligible.groupby("marketplace")[detail_metrics].sum()
    for _, row in marketplace_detail.iterrows():
        marketplace = str(row["marketplace"])
        if marketplace in totals_by_marketplace.index:
            marketplace_totals = totals_by_marketplace.loc[marketplace]
        else:
            marketplace_totals = pd.Series(0.0, index=detail_metrics)
        actual = float(marketplace_totals["paid_amount"])
        expected = float(row["net_gmv"])
        difference = actual - expected
        audit.append(
            _audit_row(
                "PASS" if difference == 0 else "FAIL",
                f"Daftar Pesanan = Net GMV {marketplace}",
                "SUM Harga Terbayarkan eligible harus sama persis dengan Net GMV.",
                actual=f"Rp {actual:,.0f}",
                expected=f"Rp {expected:,.0f}",
            )
        )

        component_checks = [
            ("seller_rebate", "Seller Rebate"),
            ("shopee_discount", "Diskon Shopee"),
            ("platform_discount", "Platform Discount"),
        ]
        for internal_column, display_label in component_checks:
            component_actual = float(marketplace_totals[internal_column])
            component_expected = float(row[internal_column])
            component_difference = component_actual - component_expected
            audit.append(
                _audit_row(
                    "PASS" if component_difference == 0 else "FAIL",
                    f"Daftar Pesanan {display_label} = MFR {marketplace}",
                    f"SUM {display_label} eligible harus sama persis dengan MFR.",
                    actual=f"Rp {component_actual:,.0f}",
                    expected=f"Rp {component_expected:,.0f}",
                )
            )

    sequence_ok = True
    for _, group in order_detail.groupby(["marketplace", "order_id"], sort=False):
        expected_sequence = list(range(1, len(group) + 1))
        if group["sequence"].astype(int).tolist() != expected_sequence:
            sequence_ok = False
            break
    audit.append(
        _audit_row(
            "PASS" if sequence_ok else "FAIL",
            "Urutan Daftar Pesanan",
            "Urutan harus reset per marketplace dan Nomor Pesanan, lalu kontinu 1..N.",
            actual="Sesuai" if sequence_ok else "Tidak sesuai",
            expected="Sesuai",
        )
    )

    shopee_nonfirst = eligible[
        eligible["marketplace"].eq("SHO") & eligible["sequence"].ne(1)
    ]
    shopee_component_once = (
        shopee_nonfirst[["seller_rebate", "shopee_discount"]]
        .fillna(0)
        .eq(0)
        .all()
        .all()
    )
    audit.append(
        _audit_row(
            "PASS" if shopee_component_once else "FAIL",
            "Komponen Order-Level Shopee Tidak Duplikat",
            "Seller Rebate dan Diskon Shopee hanya boleh muncul pada Urutan 1.",
            actual="Sesuai" if shopee_component_once else "Terduplikasi",
            expected="Sesuai",
        )
    )
    return audit


# ============================================================
# 6. EXCEL REPORT EXPORT FOR MFR
# ============================================================

def create_mfr_excel_report(mfr_result: dict[str, Any]) -> io.BytesIO:
    """
    Membuat file Excel resmi laporan MFR.
    File name contoh: MFR_Report_ID_September_2026.xlsx
    """
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "MFR Report"

    currency_format = '"Rp "#,##0'

    # Styles
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="BDD7EE")
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    bold = Font(bold=True)
    bold_title = Font(bold=True, size=13)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    row = 1

    def write_section(title: str):
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = bold_title
        ws.cell(row=row, column=1).fill = section_fill
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        row += 2

    # 1. Title & Summary
    write_section(f"LAPORAN MFR — {mfr_result.get('month_label', '')}")

    summary_items = [
        ("Periode MFR", mfr_result.get("month_label", "")),
        ("Total Net GMV", mfr_result.get("mfr_total", 0.0)),
        ("Jumlah Marketplace", mfr_result.get("marketplace_count", 0)),
        ("Total Rows Digunakan", mfr_result.get("total_rows", 0)),
    ]

    for label, val in summary_items:
        ws.cell(row=row, column=1, value=label).font = bold
        c2 = ws.cell(row=row, column=2, value=val)
        if isinstance(val, (int, float)) and label == "Total Net GMV":
            c2.number_format = currency_format
            c2.font = bold
        row += 1

    row += 2

    # 2. Ringkasan Marketplace
    write_section("RINGKASAN MFR PER MARKETPLACE")

    headers_summary = [MFR_DISPLAY_LABELS[column] for column in ["marketplace", *MFR_METRIC_COLUMNS]]
    for c_idx, h in enumerate(headers_summary, start=1):
        cell = ws.cell(row=row, column=c_idx, value=h)
        cell.font = bold
        cell.fill = header_fill
        cell.border = thin_border
    row += 1

    summary_df = mfr_result.get("marketplace_summary", pd.DataFrame()).copy()
    _require_internal_schema(
        summary_df,
        ["marketplace", *MFR_METRIC_COLUMNS],
        "ringkasan laporan Excel",
    )
    summary_display = summary_df.rename(columns=MFR_DISPLAY_LABELS)
    for _, r in summary_display.iterrows():
        is_total = r["Marketplace"] == "TOTAL MFR"
        for c_idx, h in enumerate(headers_summary, start=1):
            val = r[h]
            cell = ws.cell(row=row, column=c_idx, value=val)
            cell.border = thin_border
            if is_total:
                cell.font = bold
                cell.fill = total_fill
            if c_idx > 1 and isinstance(val, (int, float)):
                cell.number_format = currency_format
        row += 1

    row += 2

    # 3. Detail Perhitungan Marketplace
    write_section("DETAIL PERHITUNGAN MARKETPLACE")

    detail_df = mfr_result.get("marketplace_detail", pd.DataFrame())
    _require_internal_schema(
        detail_df,
        [
            "marketplace",
            *MFR_METRIC_COLUMNS,
            "source_rows",
            "included_rows",
            "excluded_rows",
        ],
        "detail laporan Excel",
    )
    headers_detail = [
        "Marketplace",
        "GMV ALL",
        "GMV Exc Batal",
        "Seller Rebate",
        "Diskon Shopee",
        "Platform Discount",
        "Net GMV",
        "Rows Total",
        "Rows Included",
        "Rows Excluded",
    ]
    for c_idx, h in enumerate(headers_detail, start=1):
        cell = ws.cell(row=row, column=c_idx, value=h)
        cell.font = bold
        cell.fill = header_fill
        cell.border = thin_border
    row += 1

    for _, r in detail_df.iterrows():
        ws.cell(row=row, column=1, value=r["marketplace"]).border = thin_border
        for c_idx, k in enumerate(MFR_METRIC_COLUMNS, start=2):
            c = ws.cell(row=row, column=c_idx, value=float(r[k]))
            c.number_format = currency_format
            c.border = thin_border
        for c_idx, k in enumerate(["source_rows", "included_rows", "excluded_rows"], start=8):
            c = ws.cell(row=row, column=c_idx, value=int(r[k]))
            c.border = thin_border
        row += 1

    row += 2

    # 4. Breakdown Status Pesanan
    write_section("BREAKDOWN STATUS PESANAN")

    status_df = mfr_result.get("status_breakdown", pd.DataFrame())
    if not status_df.empty:
        _require_internal_schema(
            status_df,
            ["marketplace", "status", "rows", "net_gmv", "treatment"],
            "breakdown status laporan Excel",
        )
    headers_status = ["Marketplace", "Status Pesanan", "Rows", "Net GMV", "Treatment"]
    for c_idx, h in enumerate(headers_status, start=1):
        cell = ws.cell(row=row, column=c_idx, value=h)
        cell.font = bold
        cell.fill = header_fill
        cell.border = thin_border
    row += 1

    for _, r in status_df.iterrows():
        ws.cell(row=row, column=1, value=r["marketplace"]).border = thin_border
        ws.cell(row=row, column=2, value=str(r["status"])).border = thin_border
        c3 = ws.cell(row=row, column=3, value=int(r["rows"]))
        c3.border = thin_border
        c4 = ws.cell(row=row, column=4, value=float(r["net_gmv"]))
        c4.number_format = currency_format
        c4.border = thin_border
        c5 = ws.cell(row=row, column=5, value=r["treatment"])
        c5.border = thin_border
        if r["treatment"] == "Excluded":
            c5.font = Font(color="9C0006")
        else:
            c5.font = Font(color="006100")
        row += 1

    row += 2

    # 5. Zero Mistake Audit
    audit_list = mfr_result.get("audit", [])
    if audit_list:
        write_section("ZERO MISTAKE RECONCILIATION AUDIT")
        headers_audit = ["Status", "Pemeriksaan", "Detail", "Actual", "Expected"]
        for c_idx, h in enumerate(headers_audit, start=1):
            cell = ws.cell(row=row, column=c_idx, value=h)
            cell.font = bold
            cell.fill = header_fill
            cell.border = thin_border
        row += 1

        for a in audit_list:
            st_val = a.get("Status", "")
            c1 = ws.cell(row=row, column=1, value=st_val)
            c1.border = thin_border
            c1.font = bold
            if st_val == "PASS":
                c1.fill = PatternFill("solid", fgColor="C6EFCE")
            elif st_val == "FAIL":
                c1.fill = PatternFill("solid", fgColor="FFC7CE")
            else:
                c1.fill = PatternFill("solid", fgColor="FFEB9C")

            ws.cell(row=row, column=2, value=str(a.get("Check", ""))).border = thin_border
            ws.cell(row=row, column=3, value=str(a.get("Detail", ""))).border = thin_border
            ws.cell(row=row, column=4, value=str(a.get("Actual", ""))).border = thin_border
            ws.cell(row=row, column=5, value=str(a.get("Expected", ""))).border = thin_border
            row += 1

    # 6. Daftar Pesanan line-item sheet
    order_detail = mfr_result.get("order_detail", pd.DataFrame()).copy()
    _require_internal_schema(
        order_detail,
        MFR_ORDER_DETAIL_COLUMNS,
        "Daftar Pesanan laporan Excel",
    )
    ws_orders = wb.create_sheet("Daftar Pesanan")
    export_columns = [
        "order_date",
        "marketplace",
        "order_id",
        "sku",
        "sequence",
        "paid_amount",
        "seller_rebate",
        "shopee_discount",
        "platform_discount",
        "order_status",
        "cancellation_reason",
    ]
    export_headers = [
        MFR_ORDER_DETAIL_DISPLAY_LABELS[column] for column in export_columns
    ]
    for column_index, header in enumerate(export_headers, start=1):
        cell = ws_orders.cell(row=1, column=column_index, value=header)
        cell.font = bold
        cell.fill = header_fill
        cell.border = thin_border

    for row_index, (_, detail_row) in enumerate(
        order_detail.iterrows(), start=2
    ):
        date_value = pd.Timestamp(detail_row["order_date"])
        date_cell = ws_orders.cell(
            row=row_index,
            column=1,
            value=date_value.to_pydatetime(),
        )
        date_cell.number_format = "d mmm yyyy"
        date_cell.border = thin_border

        values = [
            str(detail_row["marketplace"]),
            str(detail_row["order_id"]),
            str(detail_row["sku"]),
            int(detail_row["sequence"]),
            float(detail_row["paid_amount"]),
            float(detail_row["seller_rebate"]),
            float(detail_row["shopee_discount"]),
            float(detail_row["platform_discount"]),
            str(detail_row["order_status"]),
            str(detail_row["cancellation_reason"]),
        ]
        for column_index, value in enumerate(values, start=2):
            cell = ws_orders.cell(
                row=row_index,
                column=column_index,
                value=value,
            )
            cell.border = thin_border
            if column_index in {3, 4}:
                cell.number_format = "@"
            elif column_index in {6, 7, 8, 9}:
                cell.number_format = currency_format

    ws_orders.freeze_panes = "A2"
    ws_orders.auto_filter.ref = f"A1:K{max(len(order_detail) + 1, 1)}"
    detail_widths = [14, 18, 24, 18, 10, 20, 18, 18, 20, 28, 28]
    for column_index, width in enumerate(detail_widths, start=1):
        ws_orders.column_dimensions[get_column_letter(column_index)].width = width

    # Auto-adjust column widths
    for col_idx, col in enumerate(ws.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and len(val_str) < 60:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(output)
    output.seek(0)
    return output
