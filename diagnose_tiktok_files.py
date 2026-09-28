from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


BASE = Path(".")
FILES = sorted(
    p for p in BASE.glob("*.xlsx")
    if "TIK" in p.name.upper()
)


def col_to_num(col: str) -> int:
    n = 0
    for ch in col:
        if ch.isalpha():
            n = n * 26 + (ord(ch.upper()) - 64)
    return n


def inspect_xlsx(path: Path) -> None:
    print("=" * 100)
    print("FILE:", path.name)

    # 1) openpyxl view
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        print("\n[OPENPYXL]")
        print("Sheets:", wb.sheetnames)

        for ws in wb.worksheets:
            print(
                f"- {ws.title}: max_row={ws.max_row}, max_column={ws.max_column}"
            )

            for r_idx, row in enumerate(
                ws.iter_rows(
                    min_row=1,
                    max_row=min(4, ws.max_row),
                    values_only=True,
                ),
                start=1,
            ):
                vals = [
                    v
                    for v in row
                    if v not in (None, "")
                ]
                print(
                    f"  row {r_idx}: {vals[:15]}"
                )

    except Exception as exc:
        print("openpyxl ERROR:", repr(exc))

    # 2) Raw XLSX XML view
    print("\n[RAW XLSX XML]")

    try:
        with zipfile.ZipFile(path) as zf:
            worksheet_files = sorted(
                n
                for n in zf.namelist()
                if n.startswith("xl/worksheets/sheet")
                and n.endswith(".xml")
            )

            print("Worksheet XML files:", worksheet_files)

            ns = {
                "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            }

            for xml_name in worksheet_files:
                raw = zf.read(xml_name)
                root = ET.fromstring(raw)

                dim = root.find("main:dimension", ns)
                dimension = (
                    dim.attrib.get("ref")
                    if dim is not None
                    else None
                )

                refs = []
                max_col_num = 0
                max_col_letters = "A"

                for cell in root.findall(".//main:c", ns):
                    ref = cell.attrib.get("r", "")
                    if ref:
                        refs.append(ref)
                        m = re.match(r"([A-Z]+)", ref)
                        if m:
                            letters = m.group(1)
                            num = col_to_num(letters)
                            if num > max_col_num:
                                max_col_num = num
                                max_col_letters = letters

                print(
                    f"- {xml_name}: dimension={dimension}, "
                    f"highest_cell_column={max_col_letters} ({max_col_num}), "
                    f"sample_refs={refs[:20]}"
                )

    except Exception as exc:
        print("ZIP/XML ERROR:", repr(exc))

    print()


if not FILES:
    print(
        "Tidak ada file .xlsx dengan 'TIK' di nama file "
        "pada folder saat ini."
    )
else:
    print(f"Menemukan {len(FILES)} file TikTok.\n")
    for file in FILES:
        inspect_xlsx(file)
